"""Create the practical report using the bundled document runtime."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Sylfaen'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Sylfaen')
    style._element.get_or_add_rPr().rFonts.set(qn('w:cs'), 'Sylfaen')
normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(7)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12.5),('Heading 3',12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(13)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def text(value):
    return re.sub(r'\*\*|`', '', value)

def table(rows):
    if rows[0][0] == 'Հավաքածու' and 'Ընտրված մոդել' in rows[0]:
        rows = [[r[0]] + r[2:] for r in rows]
    if len(rows[0]) == 6:
        rows[0] = ['Հավաքածու', 'Հում', 'Հակասող', 'Կրկնվող', 'Մնացած', 'Train / Valid / Test']
    n = len(rows[0])
    widths = {3:[3.65,1.5,1.55],5:[1.45,1.1,1.5,1.35,1.3],6:[1.2,.9,1.0,1.0,.9,1.7]}.get(n,[6.7/n]*n)
    if rows[0][0] == "Ռեժիմ": widths = [2.05,1.05,1.2,1.2,1.2]
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for row_index,values in enumerate(rows):
        cells=t.add_row().cells
        for i,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.05
            if rows[0][0] == "Ռեժիմ": p.paragraph_format.keep_with_next = row_index < len(rows)-1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(text(value).replace('logistic_regression','Logistic Regression').replace('complement_nb','Complement NB').replace('linear_svm','Linear SVM'))
            r.font.size=Pt(10.5);r.bold=row_index==0
            tc=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            tc.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'75');m.set(qn('w:type'),'dxa');margins.append(m)
            tc.append(margins)
            if row_index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE8EF');tc.append(shade)
        tr=t.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');tr.append(no_split)
        if row_index==0: tr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=(ROOT/'PRACTICAL_REPORT_HY.md').read_text().splitlines()
rows=[]; figure_count=0
for line in lines:
    if line.startswith('|'):
        if not re.match(r'^\|[- :|]+\|$',line): rows.append([c.strip() for c in line.strip('|').split('|')])
        continue
    if rows:
        table(rows);rows=[]
    if not line: continue
    if line.startswith('# '):
        doc.add_paragraph('Կեղծ լուրերի հայտնաբերման և ճշգրտման ավտոմատացված համակարգի մշակում',style='Title')
        doc.add_paragraph('Գործնական մաս և փորձարարական արդյունքներ',style='Subtitle')
        doc.add_paragraph('Հետազոտական իրականացում անգլերեն տվյալներով',style='Subtitle')
        continue
    if line.startswith('## '):
        heading=line[3:].replace('.','').replace(':','')
        doc.add_heading(heading,level=1);continue
    if line.startswith('### '):
        doc.add_heading(line[4:].replace(':',''),level=2);continue
    if line.startswith('!['):
        path=re.search(r'\((.+)\)',line).group(1)
        figure_count+=1
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ROOT/path),width=Inches(4.7))
        p.paragraph_format.keep_with_next=True
        cap=doc.add_paragraph(f'Նկար {figure_count}  '+('LIAR' if 'liar/' in path else ('Ապացույցներով ստուգում' if 'verification/' in path else 'FakeNewsNet'))+' թեստային շփոթության մատրից',style='Caption')
        cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('F1 ='):
        doc.add_paragraph('F1-ը precision-ի և recall-ի հարմոնիկ միջինն է։ Macro-F1-ը դասերի F1-երի չկշռված միջինն է և ընտրվել է որպես հիմնական չափանիշ՝ դասերի անհավասարակշռությունը հաշվի առնելու համար։')
        continue
    doc.add_paragraph(text(line))
if rows: table(rows)
p = doc.add_heading('11 Համակարգի գործարկում և վերարտադրում',level=1)
p.paragraph_format.page_break_before = False
doc.add_paragraph('Ծրագիրը գտնվում է FakeNewsNet-master պանակի practical ենթապանակում։ start.command ֆայլը գործարկելուց հետո միջերեսը հասանելի է http://127.0.0.1:8000 հասցեով, API նկարագրությունը՝ /docs հասցեով։ Տեքստի վերլուծությունը կատարվում է /analyze մեթոդով։ Ապացույցներով ստուգումը կատարվում է /verify մեթոդով։ Առցանց MediaWiki որոնումը միացվում է համապատասխան ընտրանքով։')
doc.add_paragraph('Ուսուցման վերարտադրման համար գործարկվում է python -m fakenews.train հրամանը։ Ապացույցների բազան կառուցվում է python -m fakenews.build_evidence, իսկ որոնումը գնահատվում է python -m fakenews.evaluate_evidence հրամանով։ NLI մոդելը և սկզբնական աղբյուրները տեղադրվում են python -m fakenews.setup_verifier հրամանով, իսկ նոր շղթան գնահատվում է python -m fakenews.evaluate_verifier --split test հրամանով։ Անհրաժեշտ փաթեթների տարբերակները ֆիքսված են requirements.txt ֆայլում։')
doc.add_paragraph('Պահպանված են տվյալների մաքրման աուդիտը, train/valid/test մանիֆեստը, validation համեմատությունները, թեստային կանխատեսումները, սխալների ցանկը, շփոթության մատրիցները և աղբյուրների SHA-256 հեշերը։ Ծրագրային 60 ստուգումներն անցել են։ Բրաուզերում ստուգվել են աղբյուրով ուղղումը, փոփոխված պնդման դեպքում ձեռնպահ մնալը, արխիվային արդյունքը և փոքր էկրանի դասավորությունը։')
doc.add_paragraph('Բարելավման վերարտադրման շղթան՝ python -m fakenews.setup_encoder, python -m fakenews.improve, python -m fakenews.optimize, python -m fakenews.evaluate_optimized և python -m fakenews.activate_optimized։ Սկզբնական դասակարգիչները և նոր ընտրված տարբերակները պահպանված են առանձին։ Ծառայության վերագործարկումից հետո բեռնվում է ակտիվ գրանցամատյանը։')
doc.add_heading('12 Օգտագործված աղբյուրներ',level=1)
refs=[
'SciFact պաշտոնական հավաքածու և կառուցվածք https://github.com/allenai/scifact',
'Wang W Y 2017  Liar Liar Pants on Fire A New Benchmark Dataset for Fake News Detection  https://aclanthology.org/P17-2067/',
'Alhindi T Petridis S Muresan S 2018  Where is Your Evidence Improving Fact Checking by Justification Modeling  https://aclanthology.org/W18-5513/',
'FakeNewsNet տվյալների և հավաքման կոդի պաշտոնական պահոց  https://github.com/KaiDMML/FakeNewsNet',
'LIAR-PLUS պաշտոնական պահոց  https://github.com/Tariq60/LIAR-PLUS  commit 992bfd97555631e48b6bddb1a87aa61c51ee1197',
'scikit-learn TF-IDF փաստաթղթեր  https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html',
'Google Fact Check Tools API որոնման փաստաթղթեր  https://developers.google.com/fact-check/tools/api/reference/rest/v1alpha1/claims/search',
'NASA Sun Facts  https://science.nasa.gov/sun/facts/',
'NASA Earth Facts  https://science.nasa.gov/earth/facts/',
'Cross Encoder NLI DeBERTa v3 small մոդելի փաստաթուղթ  https://huggingface.co/cross-encoder/nli-deberta-v3-small',
'MediaWiki Search API  https://www.mediawiki.org/wiki/API:Search',
'Sentence Transformers all MiniLM L6 v2 մոդելի նկարագրություն  https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2',
]
for i,ref in enumerate(refs,1):
    p=doc.add_paragraph(f'[{i}] {ref}');p.paragraph_format.space_after=Pt(8)
    for run in p.runs: run.font.size=Pt(10.5)
doc.core_properties.title='Կեղծ լուրերի հայտնաբերման և ճշգրտման համակարգի գործնական մաս'
doc.core_properties.subject='LIAR FakeNewsNet LIAR-PLUS'
doc.core_properties.author=''
output=ROOT/'deliverables/Practical_Report_HY.docx'
# Remove inherited title rules and theme borders from every paragraph style.
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(output)
print(output)
